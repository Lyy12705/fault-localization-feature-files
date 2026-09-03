from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "reports" / "fault_localization" / "STAGE3_SYMBOL_RERANKER_IMPLEMENTATION_PLAN_ZH.docx"
EXPECTED_WIDTH = 9360
EXPECTED_INDENT = 120


def val(node, name: str) -> str:
    return node.get(qn(name)) if node is not None else ""


def main() -> None:
    errors: list[str] = []
    with ZipFile(DOCX) as archive:
        bad = archive.testzip()
        if bad:
            errors.append(f"ZIP CRC failure: {bad}")
        required = {"word/document.xml", "word/styles.xml", "word/numbering.xml"}
        missing = required.difference(archive.namelist())
        if missing:
            errors.append(f"Missing OOXML parts: {sorted(missing)}")

    doc = Document(DOCX)
    all_text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        all_text += "\n" + "\n".join(cell.text for row in table.rows for cell in row.cells)

    for token in ("TODO", "TBD", "{{", "}}", ":codex-file-citation"):
        if token in all_text:
            errors.append(f"Placeholder/internal token found: {token}")

    heading_counts = {name: 0 for name in ("Heading 1", "Heading 2", "Heading 3")}
    for paragraph in doc.paragraphs:
        if paragraph.style.name in heading_counts:
            heading_counts[paragraph.style.name] += 1
    if heading_counts["Heading 1"] < 10:
        errors.append(f"Too few Heading 1 paragraphs: {heading_counts}")

    for style_name, size, before, after in (
        ("Heading 1", 16, 360, 200),
        ("Heading 2", 13, 280, 140),
        ("Heading 3", 12, 200, 100),
    ):
        style = doc.styles[style_name]
        if round(style.font.size.pt, 2) != size:
            errors.append(f"{style_name} size mismatch: {style.font.size.pt}")
        ppr = style._element.pPr
        spacing = ppr.find(qn("w:spacing")) if ppr is not None else None
        if val(spacing, "w:before") != str(before) or val(spacing, "w:after") != str(after):
            errors.append(f"{style_name} spacing mismatch")

    normal = doc.styles["Normal"]
    normal_spacing = normal._element.pPr.find(qn("w:spacing"))
    if val(normal_spacing, "w:after") != "120" or val(normal_spacing, "w:line") != "300":
        errors.append("Normal spacing does not match compact_reference_guide")

    for t_index, table in enumerate(doc.tables, start=1):
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        grid = [int(val(col, "w:w")) for col in table._tbl.tblGrid]
        if val(tbl_w, "w:w") != str(EXPECTED_WIDTH) or val(tbl_w, "w:type") != "dxa":
            errors.append(f"Table {t_index}: tblW mismatch")
        if val(tbl_ind, "w:w") != str(EXPECTED_INDENT) or val(tbl_ind, "w:type") != "dxa":
            errors.append(f"Table {t_index}: tblInd mismatch")
        if sum(grid) != EXPECTED_WIDTH:
            errors.append(f"Table {t_index}: grid total {sum(grid)}")
        for r_index, row in enumerate(table.rows, start=1):
            widths = [int(val(cell._tc.tcPr.find(qn("w:tcW")), "w:w")) for cell in row.cells]
            if widths != grid:
                errors.append(f"Table {t_index} row {r_index}: tcW {widths} != grid {grid}")

    numbered = 0
    for paragraph in doc.paragraphs:
        ppr = paragraph._p.pPr
        if ppr is not None and ppr.find(qn("w:numPr")) is not None:
            numbered += 1
    if numbered < 20:
        errors.append(f"Expected real numbered/bulleted paragraphs, found {numbered}")

    print(
        {
            "path": str(DOCX),
            "paragraphs": len(doc.paragraphs),
            "tables": len(doc.tables),
            "headings": heading_counts,
            "numbered_paragraphs": numbered,
            "characters": len(all_text),
            "errors": errors,
        }
    )
    if errors:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
